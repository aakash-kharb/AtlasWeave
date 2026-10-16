"""Local readers. No evaluation, network access, or source-file writes."""

import csv
import re
from datetime import date, datetime, time
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook

from atlasweave.models import AtlasWeaveError, Dataset, ReadOptions, SourceAdapter

MAX_FILE_BYTES = 32 * 1024 * 1024
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_FIELDS = 100


def check_file(path: str) -> Path:
    source = Path(path)
    if not source.is_file():
        raise AtlasWeaveError(f"Source is not a file: {source}")
    if source.stat().st_size > MAX_FILE_BYTES:
        raise AtlasWeaveError("File exceeds the 32 MiB safety limit.")
    return source


def headers(values) -> tuple[str, ...]:
    result = tuple(str(value).strip() if value is not None else "" for value in values)
    if not result or len(result) > MAX_FIELDS:
        raise AtlasWeaveError("Expected between 1 and 100 columns.")
    if any(not name for name in result):
        raise AtlasWeaveError("Blank column names are unsupported; supply a complete header row.")
    if len({name.casefold() for name in result}) != len(result):
        raise AtlasWeaveError("Duplicate column names are ambiguous; rename them before importing.")
    return result


class CsvAdapter:
    storage, format, shape = "file", "csv", "table"

    def discover(self, path: str) -> list[str]:
        return [check_file(path).name]

    def read(self, path: str, options: ReadOptions) -> Dataset:
        source = check_file(path)
        if options.sheet is not None:
            raise AtlasWeaveError("Sheet selection applies only to XLSX files.")
        try:
            with source.open(encoding="utf-8-sig", newline="") as stream:
                reader = csv.reader(stream, delimiter=options.delimiter, strict=True)
                for _ in range(options.header_row - 1):
                    next(reader, None)
                names = headers(next(reader, []))
                rows = []
                complete = True
                for values in reader:
                    if not values:  # Physical blank lines, not records of empty fields.
                        continue
                    if len(values) != len(names):
                        raise AtlasWeaveError(
                            f"CSV row near line {reader.line_num} has wrong width."
                        )
                    if len(rows) == options.max_rows:
                        complete = False
                        break
                    rows.append(tuple(value if value != "" else None for value in values))
        except (UnicodeError, csv.Error) as exc:
            raise AtlasWeaveError("Invalid CSV: use UTF-8 and check quoting/delimiter.") from exc
        return Dataset(source.name, names, tuple(rows), complete)


def check_workbook(path: str) -> Path:
    source = check_file(path)
    try:
        with ZipFile(source) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                raise AtlasWeaveError("XLSX exceeds the expanded archive safety limit (64 MiB).")
            # Read-only openpyxl does not expose merged ranges. Reject them explicitly.
            for entry in entries:
                if entry.filename.startswith("xl/worksheets/") and entry.filename.endswith(".xml"):
                    with archive.open(entry) as stream:
                        for _, element in iterparse(stream, events=("end",)):
                            if element.tag.endswith("}mergeCell"):
                                raise AtlasWeaveError("Unmerge workbook cells before importing.")
                            element.clear()
    except BadZipFile as exc:
        raise AtlasWeaveError("Invalid XLSX archive.") from exc
    return source


def cell_text(cell) -> str | None:
    if cell.data_type == "f":
        raise AtlasWeaveError("Formula cells are unsupported; export a values-only sheet first.")
    if cell.data_type == "e":
        raise AtlasWeaveError("Excel error cells are unsupported; correct the source first.")
    value = cell.value
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (int, float)):
        if re.fullmatch(r"0+", cell.number_format) and float(value).is_integer():
            return str(int(value)).zfill(len(cell.number_format))
        return str(int(value)) if float(value).is_integer() else str(value)
    return str(value)


class XlsxAdapter:
    storage, format, shape = "file", "xlsx", "table"

    def discover(self, path: str) -> list[str]:
        workbook = load_workbook(check_workbook(path), read_only=True, keep_links=False)
        try:
            return workbook.sheetnames
        finally:
            workbook.close()

    def read(self, path: str, options: ReadOptions) -> Dataset:
        source = check_workbook(path)
        workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        try:
            if options.sheet is None and len(workbook.sheetnames) != 1:
                raise AtlasWeaveError(
                    "Workbook has multiple sheets; select one with --left-sheet "
                    "or --right-sheet. Use the sheets command to list them."
                )
            name = options.sheet if options.sheet is not None else workbook.sheetnames[0]
            if name not in workbook.sheetnames:
                raise AtlasWeaveError(f"Unknown worksheet: {name!r}")
            sheet = workbook[name]
            # Ignore unreliable dimension metadata; read actual worksheet cells.
            sheet.reset_dimensions()
            iterator = sheet.iter_rows(min_row=options.header_row)
            first = list(next(iterator, ()))
            while first and first[-1].value is None:
                first.pop()
            names = headers([cell_text(cell) for cell in first])
            rows = []
            complete = True
            for cells in iterator:
                if all(cell.value is None for cell in cells):
                    continue
                if len(cells) > len(names) and any(
                    c.value is not None for c in cells[len(names) :]
                ):
                    raise AtlasWeaveError("Worksheet data extends beyond the selected headers.")
                if len(rows) == options.max_rows:
                    complete = False
                    break
                values = [cell_text(cell) for cell in cells[: len(names)]]
                values.extend([None] * (len(names) - len(values)))
                rows.append(tuple(values))
            warnings = (
                "Hidden rows are included. Only simple zero-padding number formats are preserved; "
                "other numeric display formats use underlying values.",
            )
            return Dataset(source.name, names, tuple(rows), complete, warnings)
        finally:
            workbook.close()


ADAPTERS: dict[str, SourceAdapter] = {".csv": CsvAdapter(), ".xlsx": XlsxAdapter()}


def adapter_for(path: str) -> SourceAdapter:
    try:
        return ADAPTERS[Path(path).suffix.lower()]
    except KeyError as exc:
        raise AtlasWeaveError(
            "Supported file formats: .csv and .xlsx (not .xls or .xlsm)."
        ) from exc
