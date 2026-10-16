from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from atlasweave.adapters import CsvAdapter, XlsxAdapter, adapter_for
from atlasweave.cli import main
from atlasweave.models import AtlasWeaveError, ReadOptions


def test_csv_preserves_identifiers_and_quoted_multiline(tmp_path):
    path = tmp_path / "input.csv"
    path.write_text('\ufeffid;note\r\n001;"first\nsecond"\r\n002;"a;b"\r\n', encoding="utf-8")
    result = CsvAdapter().read(str(path), ReadOptions(delimiter=";"))
    assert result.fields == ("id", "note")
    assert result.rows == (("001", "first\nsecond"), ("002", "a;b"))
    assert result.complete


@pytest.mark.parametrize(
    "content, message",
    [
        ("id,ID\na,b\n", "Duplicate"),
        ("id,\na,b\n", "Blank"),
        ("id,name\na\n", "width"),
        ('id\n"unclosed\n', "Invalid CSV"),
        ("", "columns"),
    ],
)
def test_csv_invalid_input(tmp_path, content, message):
    path = tmp_path / "input.csv"
    path.write_text(content)
    with pytest.raises(AtlasWeaveError, match=message):
        CsvAdapter().read(str(path), ReadOptions())


def test_row_limits_and_header_selection(tmp_path):
    path = tmp_path / "input.csv"
    path.write_text("report\nid\n001\n002\n")
    exact = CsvAdapter().read(str(path), ReadOptions(header_row=2, max_rows=2))
    sample = CsvAdapter().read(str(path), ReadOptions(header_row=2, max_rows=1))
    assert exact.complete
    assert not sample.complete
    assert sample.rows == (("001",),)


@pytest.fixture
def workbook_path(tmp_path):
    path = tmp_path / "customers.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Customers"
    sheet.append(["Customer ID", "name"])
    sheet.append([1, "Ada Example"])
    sheet["A2"].number_format = "000"
    sheet.append(["002", "Lin Example"])
    sheet.row_dimensions[3].hidden = True
    book.save(path)
    book.close()
    return path


def test_xlsx_zero_padding_and_hidden_rows(workbook_path):
    adapter = XlsxAdapter()
    assert adapter.discover(str(workbook_path)) == ["Customers"]
    result = adapter.read(str(workbook_path), ReadOptions())
    assert result.rows == (("001", "Ada Example"), ("002", "Lin Example"))
    assert result.warnings
    assert result.complete
    sampled = adapter.read(str(workbook_path), ReadOptions(max_rows=1))
    assert not sampled.complete


def test_xlsx_ignores_incorrect_dimensions(workbook_path, tmp_path):
    corrupt = tmp_path / "dimensions.xlsx"
    with ZipFile(workbook_path) as source, ZipFile(corrupt, "w", ZIP_DEFLATED) as target:
        for name in source.namelist():
            content = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                content = content.replace(b'ref="A1:B3"', b'ref="A1:A1"')
            target.writestr(name, content)
    result = XlsxAdapter().read(str(corrupt), ReadOptions())
    assert len(result.rows) == 2
    assert len(result.fields) == 2


@pytest.mark.parametrize(
    "mode, expected",
    [
        ("formula", "Formula"),
        ("merged", "Unmerge"),
        ("error", "error cells"),
        ("extra", "beyond"),
        ("duplicate", "Duplicate"),
    ],
)
def test_xlsx_rejects_ambiguous_cells(tmp_path, mode, expected):
    path = tmp_path / "input.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["id", "name"])
    sheet.append([1, "Test"])
    if mode == "formula":
        sheet["A2"] = "=1+1"
    elif mode == "merged":
        sheet.merge_cells("A2:B2")
    elif mode == "error":
        sheet["A2"] = "#DIV/0!"
    elif mode == "extra":
        sheet["C2"] = "unheaded"
    else:
        sheet["B1"] = "ID"
    book.save(path)
    book.close()
    with pytest.raises(AtlasWeaveError, match=expected):
        XlsxAdapter().read(str(path), ReadOptions())


def test_multiple_sheets_require_selection(tmp_path):
    path = tmp_path / "multi.xlsx"
    book = Workbook()
    book.active.append(["id"])
    book.create_sheet("Other").append(["other_id"])
    book.save(path)
    book.close()
    with pytest.raises(AtlasWeaveError, match="multiple"):
        XlsxAdapter().read(str(path), ReadOptions())
    assert XlsxAdapter().read(str(path), ReadOptions(sheet="Other")).fields == ("other_id",)
    with pytest.raises(AtlasWeaveError, match="Unknown"):
        XlsxAdapter().read(str(path), ReadOptions(sheet="Missing"))


def test_xlsx_csv_complete_workflow(workbook_path, tmp_path):
    left = tmp_path / "orders.csv"
    left.write_text("cust_id,order\n001,o1\n002,o2\n003,o3\n")
    project = tmp_path / "project.json"
    assert main(["discover", str(left), str(workbook_path), "--out", str(project)]) == 0
    assert main(["review", str(project), "left.f1--right.f1", "--status", "approved"]) == 0
    assert main(["export", str(project), "--out-dir", str(tmp_path / "export")]) == 0


def test_archive_expansion_limit(tmp_path, monkeypatch):
    from atlasweave import adapters

    path = tmp_path / "large.xlsx"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("large", b"x" * 1024)
    monkeypatch.setattr(adapters, "MAX_EXPANDED_BYTES", 100)
    with pytest.raises(AtlasWeaveError, match="expanded"):
        XlsxAdapter().discover(str(path))


@pytest.mark.parametrize(
    "xml",
    [
        b"<worksheet><broken>",
        b'<!DOCTYPE worksheet [<!ENTITY x "expanded">]><worksheet>&x;</worksheet>',
    ],
)
def test_malformed_and_entity_xml_fail_cleanly(tmp_path, xml, capsys):
    path = tmp_path / "unsafe.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr("xl/worksheets/sheet1.xml", xml)
    assert main(["sheets", str(path)]) == 2
    assert "Traceback" not in capsys.readouterr().err


@pytest.mark.parametrize(
    "options",
    [
        {"max_rows": 0},
        {"max_rows": 100001},
        {"header_row": 0},
        {"delimiter": "||"},
    ],
)
def test_invalid_read_options(options):
    with pytest.raises(AtlasWeaveError):
        ReadOptions(**options)


def test_unsupported_formats_and_invalid_xlsx(tmp_path):
    with pytest.raises(AtlasWeaveError, match="Supported"):
        adapter_for("old.xls")
    path = tmp_path / "bad.xlsx"
    path.write_text("not an archive")
    assert main(["sheets", str(path)]) == 2
